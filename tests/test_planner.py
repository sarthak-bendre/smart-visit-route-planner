import numpy as np
import pandas as pd
import pytest

from data.generate_dealers import generate
from planner import BRANCH_OFFICE
from planner.distance import ROAD_FACTOR, estimate_matrix, haversine_km
from planner.optimizer import (
    PlanningError,
    PlanSettings,
    plan_route,
    route_metrics,
    unplanned_order,
)
from planner.priority import score_dealers, tier
from planner.report import build_schedule, compute_kpis, not_visited, summary_table, to_excel


@pytest.fixture(scope="module")
def dealers():
    return score_dealers(generate())


@pytest.fixture(scope="module")
def matrix(dealers):
    points = [(BRANCH_OFFICE["latitude"], BRANCH_OFFICE["longitude"])]
    points += list(zip(dealers["latitude"], dealers["longitude"]))
    return estimate_matrix(points)


# ----------------------------------------------------------------------------- priority
def test_score_formula_matches_report():
    df = pd.DataFrame(
        {
            "monthly_target_lakh": [20.0, 40.0],
            "achieved_lakh": [10.0, 50.0],  # 50% gap; above target -> 0 gap
            "pending_issues": [2, 9],  # 9 is capped at 5
            "days_since_last_visit": [15, 60],  # 60 is capped at 30
        }
    )
    out = score_dealers(df)
    # 100 * (0.35*0.5 + 0.25*0.4 + 0.25*0.5 + 0.15*0.5) = 47.5
    assert out.loc[0, "priority_score"] == pytest.approx(47.5)
    # 100 * (0 + 0.25*1 + 0.25*1 + 0.15*1) = 65
    assert out.loc[1, "priority_score"] == pytest.approx(65.0)


def test_weights_are_normalised():
    df = pd.DataFrame({"monthly_target_lakh": [10.0], "achieved_lakh": [0.0],
                       "pending_issues": [0], "days_since_last_visit": [0]})
    out = score_dealers(df, {"target_gap": 2, "pending_issues": 0, "days_since_visit": 0, "dealer_size": 0})
    assert out.loc[0, "priority_score"] == pytest.approx(100.0)


def test_tiers():
    assert tier(40) == "High" and tier(39.9) == "Medium" and tier(30) == "Medium" and tier(29.9) == "Low"


# ----------------------------------------------------------------------------- distance
def test_haversine_known_distance():
    # Shivajinagar to Hinjewadi is roughly 14 km in a straight line.
    assert 12 < haversine_km(18.5308, 73.8475, 18.5913, 73.7389) < 15


def test_estimate_matrix_is_symmetric_with_road_factor(dealers, matrix):
    first_dealer = (dealers.loc[0, "latitude"], dealers.loc[0, "longitude"])
    assert np.allclose(matrix.km, matrix.km.T)
    assert np.all(np.diag(matrix.km) == 0)
    straight = haversine_km(BRANCH_OFFICE["latitude"], BRANCH_OFFICE["longitude"], *first_dealer)
    assert matrix.km[0, 1] == pytest.approx(ROAD_FACTOR * straight)
    assert np.allclose(matrix.minutes, matrix.km / 20 * 60)


# ----------------------------------------------------------------------------- optimiser
@pytest.mark.parametrize("use_ortools", [True, False])
def test_plan_respects_constraints(dealers, matrix, use_ortools):
    settings = PlanSettings(visit_minutes=30, day_hours=8, max_stops=10, time_limit_s=1)
    plan = plan_route(matrix.km, matrix.minutes, dealers["priority_score"], set(), settings, use_ortools)
    assert 0 < len(plan.order) <= settings.max_stops
    assert len(set(plan.order)) == len(plan.order)
    assert sorted(plan.order + plan.skipped) == list(range(len(dealers)))
    assert plan.total_minutes <= settings.day_minutes + 1e-6


@pytest.mark.parametrize("use_ortools", [True, False])
def test_must_visit_always_included(dealers, matrix, use_ortools):
    lowest = int(dealers["priority_score"].idxmin())
    settings = PlanSettings(max_stops=5, time_limit_s=1)
    plan = plan_route(matrix.km, matrix.minutes, dealers["priority_score"], {lowest}, settings, use_ortools)
    assert lowest in plan.order


@pytest.mark.parametrize("use_ortools", [True, False])
def test_low_priority_dropped_first(dealers, matrix, use_ortools):
    settings = PlanSettings(max_stops=6, time_limit_s=1, km_per_priority_point=5)
    plan = plan_route(matrix.km, matrix.minutes, dealers["priority_score"], set(), settings, use_ortools)
    visited = dealers.loc[plan.order, "priority_score"].mean()
    skipped = dealers.loc[plan.skipped, "priority_score"].mean()
    assert visited > skipped


def test_too_many_must_visits_rejected(dealers, matrix):
    with pytest.raises(PlanningError):
        plan_route(matrix.km, matrix.minutes, dealers["priority_score"], set(range(6)), PlanSettings(max_stops=5))


@pytest.mark.parametrize("use_ortools", [True, False])
def test_infeasible_must_visits_rejected(dealers, matrix, use_ortools):
    settings = PlanSettings(day_hours=1, visit_minutes=25, max_stops=5, time_limit_s=1)
    far = list(dealers["priority_score"].index[:2])
    with pytest.raises(PlanningError):
        plan_route(matrix.km, matrix.minutes, dealers["priority_score"], set(far), settings, use_ortools)


def test_planned_route_beats_unplanned(dealers, matrix):
    plan = plan_route(matrix.km, matrix.minutes, dealers["priority_score"], set(), PlanSettings(time_limit_s=1))
    base_km, _, _ = route_metrics(unplanned_order(plan.order, dealers["priority_score"].to_numpy()),
                                  matrix.km, matrix.minutes, 30)
    assert plan.total_km <= base_km


# ----------------------------------------------------------------------------- report
def test_schedule_and_excel(dealers, matrix):
    settings = PlanSettings(time_limit_s=1)
    plan = plan_route(matrix.km, matrix.minutes, dealers["priority_score"], set(), settings)
    schedule = build_schedule(dealers, plan, matrix.km, matrix.minutes, settings.visit_minutes)
    assert len(schedule) == len(plan.order) + 1
    assert schedule.iloc[-1]["Stop"] == "End"
    assert schedule["Drive km"].sum() == pytest.approx(plan.total_km, abs=0.1 * len(schedule))
    assert schedule["Arrive"].tolist() == sorted(schedule["Arrive"].tolist())  # same-day, increasing

    kpis = compute_kpis(dealers, plan, matrix.km, matrix.minutes, settings.visit_minutes)
    assert kpis["dealers_visited"] == len(plan.order)
    xlsx = to_excel(schedule, not_visited(dealers, plan), summary_table(kpis, {"Solver": plan.method}))

    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(xlsx))
    assert wb.sheetnames == ["Visit Plan", "Not Visited Today", "Summary"]
    assert wb["Visit Plan"].cell(row=2, column=11).hyperlink is not None
