"""Step 4: turn a route plan into a timed schedule, KPIs, an interactive map and an Excel report."""
from datetime import datetime, time, timedelta
from io import BytesIO

import folium
import numpy as np
import pandas as pd

from planner import BRANCH_OFFICE
from planner.optimizer import RoutePlan, route_metrics, unplanned_order

TIER_COLOURS = {"High": "#D64545", "Medium": "#E8912D", "Low": "#2E9E6A"}
SKIPPED_COLOUR = "#A0A7B4"
ROUTE_COLOUR = "#1F4E9C"
OFFICE_COLOUR = "#14233C"

# Muted basemap so the coloured stops stand out, plus popup and legend styling.
MAP_CSS = """
<style>
.leaflet-tile-pane { filter: grayscale(0.85) contrast(0.92) brightness(1.04); }
.leaflet-container { font-family: Inter, -apple-system, Segoe UI, Roboto, sans-serif; }
.leaflet-popup-content-wrapper { border-radius: 10px; box-shadow: 0 6px 24px rgba(20,35,60,.18); }
.leaflet-popup-content { margin: 12px 14px; font-size: 12.5px; line-height: 1.5; color: #1A2233; }
.pp-head { font-size: 11px; font-weight: 600; letter-spacing: .04em; text-transform: uppercase; color: #6B7485; }
.pp-name { font-size: 14px; font-weight: 700; margin: 2px 0 6px; }
.pp-row { display: flex; justify-content: space-between; gap: 16px; }
.pp-row span:first-child { color: #6B7485; }
.map-legend { position: absolute; bottom: 18px; left: 12px; z-index: 1000; background: rgba(255,255,255,.95);
  border-radius: 10px; padding: 10px 12px; box-shadow: 0 2px 12px rgba(20,35,60,.15);
  font: 12px Inter, -apple-system, Segoe UI, Roboto, sans-serif; color: #1A2233; }
.map-legend div { display: flex; align-items: center; gap: 8px; margin: 3px 0; }
.map-legend i { width: 12px; height: 12px; border-radius: 50%; display: inline-block; border: 2px solid #fff;
  box-shadow: 0 0 0 1px rgba(0,0,0,.12); }
.map-legend b { font-size: 11px; letter-spacing: .04em; text-transform: uppercase; color: #6B7485; }
</style>
"""


def maps_link(origin: tuple[float, float], dest: tuple[float, float]) -> str:
    return (
        "https://www.google.com/maps/dir/?api=1"
        f"&origin={origin[0]:.6f},{origin[1]:.6f}"
        f"&destination={dest[0]:.6f},{dest[1]:.6f}&travelmode=driving"
    )


def _office_point() -> tuple[float, float]:
    return BRANCH_OFFICE["latitude"], BRANCH_OFFICE["longitude"]


def build_schedule(
    dealers: pd.DataFrame,
    plan: RoutePlan,
    km: np.ndarray,
    minutes: np.ndarray,
    visit_minutes: float,
    start: time = time(9, 30),
) -> pd.DataFrame:
    """One row per stop with arrival/leaving times, plus the final drive back to the office."""
    clock = datetime.combine(datetime.today(), start)
    prev_node, prev_point = 0, _office_point()
    rows = []
    for stop, dealer in enumerate(plan.order, start=1):
        node = dealer + 1
        d = dealers.iloc[dealer]
        point = (float(d["latitude"]), float(d["longitude"]))
        arrive = clock + timedelta(minutes=float(minutes[prev_node, node]))
        leave = arrive + timedelta(minutes=visit_minutes)
        rows.append(
            {
                "Stop": str(stop),
                "Dealer ID": d["dealer_id"],
                "Dealer": d["dealer_name"],
                "Area": d["area"],
                "Priority": float(d["priority_score"]),
                "Tier": d["priority_tier"],
                "Drive km": round(float(km[prev_node, node]), 1),
                "Drive min": round(float(minutes[prev_node, node])),
                "Arrive": arrive.strftime("%H:%M"),
                "Leave": leave.strftime("%H:%M"),
                "Navigate": maps_link(prev_point, point),
            }
        )
        clock, prev_node, prev_point = leave, node, point

    back = clock + timedelta(minutes=float(minutes[prev_node, 0]))
    rows.append(
        {
            "Stop": "End",
            "Dealer ID": "",
            "Dealer": BRANCH_OFFICE["name"],
            "Area": "",
            "Priority": None,
            "Tier": "",
            "Drive km": round(float(km[prev_node, 0]), 1),
            "Drive min": round(float(minutes[prev_node, 0])),
            "Arrive": back.strftime("%H:%M"),
            "Leave": "",
            "Navigate": maps_link(prev_point, _office_point()),
        }
    )
    return pd.DataFrame(rows)


def not_visited(dealers: pd.DataFrame, plan: RoutePlan) -> pd.DataFrame:
    cols = ["dealer_id", "dealer_name", "area", "priority_score", "priority_tier",
            "pending_issues", "days_since_last_visit"]
    out = dealers.iloc[plan.skipped][cols].sort_values("priority_score", ascending=False)
    return out.rename(columns={
        "dealer_id": "Dealer ID", "dealer_name": "Dealer", "area": "Area",
        "priority_score": "Priority", "priority_tier": "Tier",
        "pending_issues": "Pending issues", "days_since_last_visit": "Days since visit",
    }).reset_index(drop=True)


def compute_kpis(dealers: pd.DataFrame, plan: RoutePlan, km: np.ndarray, minutes: np.ndarray, visit_minutes: float) -> dict:
    priority = dealers["priority_score"].to_numpy()
    base_km, base_drive, base_total = route_metrics(
        unplanned_order(plan.order, priority), km, minutes, visit_minutes
    )
    visited = dealers.iloc[plan.order]
    high_total = int((dealers["priority_tier"] == "High").sum())
    km_saved = base_km - plan.total_km
    return {
        "dealers_visited": len(plan.order),
        "dealers_total": len(dealers),
        "high_visited": int((visited["priority_tier"] == "High").sum()),
        "high_total": high_total,
        "planned_km": round(plan.total_km, 1),
        "drive_hours": round(plan.travel_minutes / 60, 2),
        "day_hours": round(plan.total_minutes / 60, 2),
        "unplanned_km": round(base_km, 1),
        "unplanned_day_hours": round(base_total / 60, 2),
        "km_saved": round(km_saved, 1),
        "pct_saved": round(100 * km_saved / base_km, 1) if base_km > 0 else 0.0,
        "minutes_saved": round(base_drive - plan.travel_minutes),
        "priority_covered": round(float(visited["priority_score"].sum()), 1),
    }


def _popup(d: pd.Series, header: str) -> folium.Popup:
    from html import escape

    rows = [
        ("Priority", f"<b>{d['priority_score']:.1f}</b> &middot; {d['priority_tier']}"),
        ("Achieved", f"{d['achieved_lakh']:.1f} / {d['monthly_target_lakh']:.0f} lakh"),
        ("Pending issues", d["pending_issues"]),
        ("Last visit", f"{d['days_since_last_visit']} days ago"),
    ]
    body = "".join(f'<div class="pp-row"><span>{k}</span><span>{v}</span></div>' for k, v in rows)
    return folium.Popup(
        f'<div class="pp-head">{escape(header)}</div>'
        f'<div class="pp-name">{escape(str(d["dealer_name"]))}</div>'
        f'<div style="color:#6B7485;margin-bottom:6px">{escape(str(d["area"]))} &middot; {escape(str(d["dealer_id"]))}</div>'
        f"{body}",
        max_width=260,
    )


def _legend_html() -> str:
    items = "".join(f'<div><i style="background:{c}"></i>{t} priority</div>' for t, c in TIER_COLOURS.items())
    return (
        f'<div class="map-legend"><b>Legend</b>{items}'
        f'<div><i style="background:{SKIPPED_COLOUR}"></i>Not visited today</div>'
        f'<div><i style="background:{OFFICE_COLOUR};border-radius:3px"></i>Branch office</div></div>'
    )


def build_map(dealers: pd.DataFrame, plan: RoutePlan, geometry: list[tuple[float, float]]) -> folium.Map:
    office = _office_point()
    fmap = folium.Map(location=office, zoom_start=12, tiles="OpenStreetMap", fade_animation=False, zoom_snap=0.25)
    fmap.get_root().header.add_child(folium.Element(MAP_CSS))
    fmap.get_root().html.add_child(folium.Element(_legend_html()))

    if len(geometry) >= 2:
        # White casing under the route line makes it readable over any basemap.
        folium.PolyLine(geometry, color="#FFFFFF", weight=9, opacity=0.9).add_to(fmap)
        folium.PolyLine(geometry, color=ROUTE_COLOUR, weight=5, opacity=0.9, tooltip="Planned route").add_to(fmap)

    for dealer in plan.skipped:
        d = dealers.iloc[dealer]
        folium.CircleMarker(
            (d["latitude"], d["longitude"]), radius=6, color="#FFFFFF", fill=True,
            fill_color=SKIPPED_COLOUR, fill_opacity=0.95, weight=2,
            tooltip=f"Not today: {d['dealer_name']} ({d['priority_score']:.0f})",
            popup=_popup(d, "Not visited today"),
        ).add_to(fmap)

    for stop, dealer in enumerate(plan.order, start=1):
        d = dealers.iloc[dealer]
        colour = TIER_COLOURS.get(d["priority_tier"], ROUTE_COLOUR)
        html = (
            f'<div style="background:{colour};color:#fff;border:2.5px solid #fff;border-radius:50%;'
            f"width:28px;height:28px;display:flex;align-items:center;justify-content:center;"
            f"font:700 13px Inter,-apple-system,Segoe UI,Roboto,sans-serif;"
            f'box-shadow:0 2px 6px rgba(20,35,60,.35);">{stop}</div>'
        )
        folium.Marker(
            (d["latitude"], d["longitude"]),
            icon=folium.DivIcon(html=html, icon_size=(28, 28), icon_anchor=(14, 14)),
            tooltip=f"{stop}. {d['dealer_name']} ({d['priority_tier']})",
            popup=_popup(d, f"Stop {stop}"),
        ).add_to(fmap)

    office_html = (
        f'<div style="background:{OFFICE_COLOUR};color:#fff;border:2.5px solid #fff;border-radius:8px;'
        f"width:34px;height:30px;display:flex;align-items:center;justify-content:center;"
        f"font:800 11px Inter,-apple-system,Segoe UI,Roboto,sans-serif;letter-spacing:.04em;"
        f'box-shadow:0 2px 8px rgba(20,35,60,.4);">HQ</div>'
    )
    folium.Marker(
        office,
        icon=folium.DivIcon(html=office_html, icon_size=(34, 30), icon_anchor=(17, 15)),
        tooltip=BRANCH_OFFICE["name"],
        popup=folium.Popup(f'<div class="pp-head">Start &amp; end</div>'
                           f'<div class="pp-name">{BRANCH_OFFICE["name"]}</div>', max_width=240),
        z_index_offset=1000,
    ).add_to(fmap)

    points = [office] + [(dealers.iloc[i]["latitude"], dealers.iloc[i]["longitude"]) for i in range(len(dealers))]
    fmap.fit_bounds([[min(p[0] for p in points), min(p[1] for p in points)],
                     [max(p[0] for p in points), max(p[1] for p in points)]], padding=(24, 24))
    return fmap


def summary_table(kpis: dict, settings_info: dict) -> pd.DataFrame:
    rows = [
        ("Dealers visited", f"{kpis['dealers_visited']} / {kpis['dealers_total']}"),
        ("High-priority dealers visited", f"{kpis['high_visited']} / {kpis['high_total']}"),
        ("Planned route (km)", kpis["planned_km"]),
        ("Unplanned route (km)", kpis["unplanned_km"]),
        ("Distance saved (km)", kpis["km_saved"]),
        ("Distance saved (%)", kpis["pct_saved"]),
        ("Drive time saved (min)", kpis["minutes_saved"]),
        ("Drive time (h)", kpis["drive_hours"]),
        ("Day length incl. visits (h)", kpis["day_hours"]),
        ("Total priority covered", kpis["priority_covered"]),
    ]
    rows += list(settings_info.items())
    return pd.DataFrame(rows, columns=["Metric", "Value"]).astype({"Value": str})


def to_excel(schedule: pd.DataFrame, skipped: pd.DataFrame, summary: pd.DataFrame) -> bytes:
    """Excel workbook with three sheets: Visit Plan, Not Visited Today, Summary."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, frame in (("Visit Plan", schedule), ("Not Visited Today", skipped), ("Summary", summary)):
            frame.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E9C")
            for idx, col in enumerate(frame.columns, start=1):
                width = max([len(str(col))] + [len(str(v)) for v in frame[col].head(50)])
                ws.column_dimensions[get_column_letter(idx)].width = min(width + 2, 45)
            ws.freeze_panes = "A2"

        # Make the navigation links clickable.
        ws = writer.sheets["Visit Plan"]
        if "Navigate" in schedule.columns:
            col = schedule.columns.get_loc("Navigate") + 1
            for row in range(2, len(schedule) + 2):
                cell = ws.cell(row=row, column=col)
                if cell.value:
                    cell.hyperlink = cell.value
                    cell.value = "Open in Google Maps"
                    cell.font = Font(color="0563C1", underline="single")
            ws.column_dimensions[get_column_letter(col)].width = 22
    return buf.getvalue()
